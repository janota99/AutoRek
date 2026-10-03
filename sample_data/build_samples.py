"""Regenerates the synthetic sample files in this folder (`py sample_data/build_samples.py`).

Everything here is invented: made-up vendors, customers, PO numbers and amounts, shaped like
the real exports so each app's "Use Sample Data" button and the downloadable templates work.
Never copy real company data into this folder; it is committed.
"""

from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------- Sales Tax
VENDORS = [  # (name, id, taxability, grouping)
    ("Sample Packaging Co", "900101", "Taxable", "Supplies"),
    ("Sample Label Works", "900102", "Taxable", "Supplies"),
    ("Sample Freight Lines", "900103", "Exempt", "Freight"),
    ("Sample Chemical Supply", "900104", "Taxable", "Chemicals"),
    ("Sample Office Depot", "900105", "Taxable", "Office"),
    ("Sample Utilities LLC", "900106", "Exempt", "Utilities"),
]
NEW_VENDORS = ["Sample Pallet Brokers", "Sample Uniform Rental"]  # not in the mapping: flagged as new
EXCLUDED_VENDOR = "101806"  # on the default exclusion list

SOURCE_COLS = ["Vendor", "Invoice Date", "Transaction ID", "Description", "Account #",
               "Cost Center", "Amount", "Tax Code", "Reference", "Code", "Period"]
GL = {  # account -> name (cost center is always 10100)
    610000: "Packaging Supplies", 620000: "Freight In", 630000: "Chemicals",
    640000: "Office Supplies", 650000: "Utilities",
}
rows = []
spec = [(v[0], a) for v, a in zip(VENDORS, GL)] + [(v[0], 610000) for v in VENDORS[:2]]
for i, (vendor, acct) in enumerate(spec):
    rows.append([vendor, f"2026-07-{i + 1:02d}", f"TX{1000 + i}", f"Invoice {i + 1}", acct, 10100,
                 round(125.5 + i * 37.25, 2), "T", f"REF{i + 1}", "", 12])
rows.append([NEW_VENDORS[0], "2026-07-20", "TX2000", "Pallets", 610000, 10100, 480.00, "T", "REF20", "", 12])
rows.append([NEW_VENDORS[1], "2026-07-21", "TX2001", "Uniforms", 640000, 10100, 215.40, "T", "REF21", "", 12])
rows.append([EXCLUDED_VENDOR, "2026-07-22", "TX2002", "Excluded vendor", 610000, 10100, 99.99, "T", "REF22", "", 12])
rows.append([VENDORS[0][0], "2026-07-23", "TX2003", "Letter-prefixed code", 610000, 10100, 50.00, "T", "REF23", "A-ADJ", 12])
rows.append([VENDORS[0][0], "2026-07-24", "TX1000", "Duplicate of TX1000", 610000, 10100, 125.50, "T", "REF24", "", 12])
pd.DataFrame(rows, columns=SOURCE_COLS).to_excel(HERE / "sales_tax_source_transactions.xlsx", index=False)

pd.DataFrame(
    [[n, i, "", t, g, "Active"] for n, i, t, g in VENDORS],
    columns=["Vendor Name", "Vendor ID", "Notes", "Taxability", "Grouping", "Status"],
).to_excel(HERE / "sales_tax_vendor_mapping.xlsx", index=False)

pd.DataFrame(
    [["TB", name, f"001-10100-{acct}-000", 0.0] for acct, name in GL.items()],
    columns=["Entity", "Account Name", "GL Account", "Balance"],
).to_excel(HERE / "sales_tax_trial_balance.xlsx", index=False)

# Vendor Reconciliation: one added, one removed, one renamed, rest unchanged
old = [[i, n, t, g] for n, i, t, g in VENDORS]
new = [[i, n, t, g] for n, i, t, g in VENDORS]
new[1][1] = "Sample Label Works Inc"       # renamed
new.pop()                                  # removed (900106)
new.append(["900107", "Sample Pallet Brokers", "Taxable", "Supplies"])  # added
cols = ["Vendor ID", "Vendor Name", "Taxability", "Grouping"]
pd.DataFrame(old, columns=cols).to_excel(HERE / "vendor_listing_previous.xlsx", index=False)
pd.DataFrame(new, columns=cols).to_excel(HERE / "vendor_listing_updated.xlsx", index=False)

# -------------------------------------------------------------------- Recon
CUSTOMERS = ["SAMPLE FOODS, INC.", "DEMO GROCERS CO", "EXAMPLE MARKETS LLC"]


def qb_frame(items):
    return pd.DataFrame(items, columns=["Date", "Num", "Memo/Description", "P.O. NUMBER",
                                        "Customer name", "QTY", "AMOUNT"])


def write_qb(path, items, period=10):
    """QuickBooks-style export: report title rows above the header, blank row below it."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "SAMPLE Sales by Product/Service Detail"
    ws["A2"] = "October 2026 (synthetic sample data)"
    ws.append([])
    ws.append([])
    # Header on row 5. Column A is the fiscal period (the Recon page reads it as the QuickBooks
    # period), B is left blank like the real export, data starts in column C.
    ws.append(["PERIOD", None, *qb_frame([]).columns])
    ws.append([])
    for item in items:
        ws.append([period, None, *item])
    wb.save(path)


qb, inf = [], []
for i in range(12):
    po, amount, qty = 800000 + i, round(2000 + i * 111.11, 2), 1000 + i * 10
    cust = CUSTOMERS[i % 3]
    qb.append([f"10/{i + 1:02d}/2026", 30000 + i, "SAMPLE 24 CASE", po, cust, qty, amount])
    inf.append([10, f"10/{i + 1:02d}/2026", "WP", 50000 + i % 3, 20000 + i, amount, f"PO {po}"])
qb.append(["10/15/2026", 30100, "SAMPLE 32 CASE", 800500, CUSTOMERS[0], 900, 1500.00])   # QB only
qb.append(["10/16/2026", 30101, "SAMPLE 40 CASE", 800501, CUSTOMERS[1], 700, 980.25])    # QB only
inf.append([10, "10/17/2026", "WP", 50001, 20100, 2750.00, "PO 800600"])                   # Infinium only
inf.append([10, "10/18/2026", "WP", 50002, 20101, 640.10, "PO 800601"])                    # Infinium only
inf[3][5] = round(inf[3][5] + 10.00, 2)                                                  # amount differs

write_qb(HERE / "recon_quickbooks_sales.xlsx", qb)
inf_cols = ["OHAPD", "OHOBDE", "OHCO", "CUNO", "OHOBNO", "OHTOTA", "OHDESC"]
pd.DataFrame(inf, columns=inf_cols).to_excel(HERE / "recon_infinium_sales.xlsx", index=False)

# Prior-period rows that clear the Infinium-only items (QB side) and the QB-only items (Infinium side)
write_qb(HERE / "recon_quickbooks_historical.xlsx", [
    ["09/28/2026", 29900, "SAMPLE 24 CASE", 800600, CUSTOMERS[2], 800, 2750.00],
    ["09/29/2026", 29901, "SAMPLE 32 CASE", 800601, CUSTOMERS[0], 300, 640.10],
], period=9)
pd.DataFrame([
    [9, "9/27/2026", "WP", 50000, 19900, 1500.00, "PO 800500"],
    [9, "9/28/2026", "WP", 50001, 19901, 980.25, "PO 800501"],
], columns=inf_cols).to_excel(HERE / "recon_infinium_historical.xlsx", index=False)
print("wrote", len(list(HERE.glob("*.xlsx"))), "files")
