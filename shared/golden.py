"""Golden-file helpers for tests: freeze today's output, fail loudly if it ever changes.

Used by the FIFO and Sales Tax tests, which prove that a refactor changed no accounting result.
A golden file is plain JSON stored next to the test. To accept an intended change, re-run with
``GOLDEN_UPDATE=1`` and review the diff in git before committing it; never update to make a
failing test pass without understanding why the numbers moved.

Test-only: nothing in the running app imports this module.
"""
from __future__ import annotations

import json
import os
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

_FLOAT_DIGITS = 9  # tolerate platform float noise far below a cent, never anything a cent could see


def _plain(value):
    """Recursively turn frames, numpy scalars, dates, and floats into JSON-stable values."""
    if isinstance(value, pd.DataFrame):
        return {
            "columns": [str(c) for c in value.columns],
            "dtypes": [str(t) for t in value.dtypes],
            "rows": [[_plain(v) for v in row] for row in value.itertuples(index=False, name=None)],
        }
    if isinstance(value, pd.Series):
        return {"values": [_plain(v) for v in value.tolist()]}
    if isinstance(value, dict):
        return {str(k): _plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_plain(v) for v in value]
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(value, float):
        return None if value != value else round(value, _FLOAT_DIGITS)
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if value is pd.NaT:
        return None
    return str(value)


def workbook_snapshot(source) -> dict:
    """Every sheet's name, dimensions, freeze pane, widths, merged ranges, and each non-empty
    cell's value, number format, bold flag, and fill colour. ``source`` is bytes, a BytesIO, or
    an openpyxl workbook."""
    if isinstance(source, (bytes, bytearray)):
        wb = load_workbook(BytesIO(source))
    elif isinstance(source, BytesIO):
        wb = load_workbook(BytesIO(source.getvalue()))
    else:
        wb = source
    sheets = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
                cells[cell.coordinate] = {
                    "v": _plain(cell.value),
                    "fmt": cell.number_format,
                    "bold": bool(cell.font and cell.font.b),
                    "fill": fill if isinstance(fill, str) else None,
                }
        sheets[ws.title] = {
            "dimensions": ws.dimensions,
            "freeze": ws.freeze_panes,
            "filter": ws.auto_filter.ref,
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "widths": {k: round(v.width, 2) for k, v in sorted(ws.column_dimensions.items()) if v.width},
            "cells": cells,
        }
    return {"sheet_order": [ws.title for ws in wb.worksheets], "sheets": sheets}


def assert_matches_golden(golden_dir: Path, name: str, actual) -> None:
    """Compare ``actual`` with ``<golden_dir>/<name>.json`` (or rewrite it under GOLDEN_UPDATE=1)."""
    path = Path(golden_dir) / f"{name}.json"
    text = json.dumps(_plain(actual), indent=1, sort_keys=True, ensure_ascii=False) + "\n"
    if os.environ.get("GOLDEN_UPDATE") == "1":
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8", newline="\n")
        return
    if not path.exists():
        raise AssertionError(f"No golden file at {path}. Create it with GOLDEN_UPDATE=1 and review it.")
    expected = path.read_text(encoding="utf-8")
    if expected != text:
        want, got = json.loads(expected), json.loads(text)
        raise AssertionError(
            f"{name}: output differs from {path.name}. First difference: {_first_difference(want, got)}"
        )


def _first_difference(want, got, where="$"):
    if type(want) is not type(got):
        return f"{where}: expected {want!r}, got {got!r}"
    if isinstance(want, dict):
        for key in sorted(set(want) | set(got)):
            if key not in want:
                return f"{where}.{key}: unexpected key"
            if key not in got:
                return f"{where}.{key}: missing key"
            found = _first_difference(want[key], got[key], f"{where}.{key}")
            if found:
                return found
        return None
    if isinstance(want, list):
        if len(want) != len(got):
            return f"{where}: expected {len(want)} items, got {len(got)}"
        for i, (a, b) in enumerate(zip(want, got)):
            found = _first_difference(a, b, f"{where}[{i}]")
            if found:
                return found
        return None
    return None if want == got else f"{where}: expected {want!r}, got {got!r}"
