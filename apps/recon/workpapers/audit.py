"""The Audit & Controls sheet: run identity, approval, export controls, and the reviewer layer."""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..config import FONT_NAME, GREEN_LIGHT, NAVY, NAVY_LIGHT, RED_LIGHT, SLATE, SLATE_LIGHT, TEXT
from ..excel_styles import (
    _apply_number_formats,
    _format_body_block,
    _format_header,
    _prepare_sheet,
    _write_caption_band,
    _write_title_band,
    fix_row_height,
)
from ..matching import APP_VERSION, MATCHING_RULE_VERSION, ReconciliationResult
from ..duplicates import DUPLICATE_RULE_VERSION
from ..matching.export_checks import LIVE_WORKBOOK_CONTROLS
from ..matching.review_decisions import adjusted_je_amount, approval_status_text
from .tables import _write_dataframe_values


AUDIT_SHEET = "Audit & Controls"

_END_COL = 6


def _write_table(ws, row: int, frame: pd.DataFrame, title: str, caption: str, color: str, amounts: set[str]) -> int:
    """A titled table; returns the next free row (one blank row after it)."""
    end_col = max(len(frame.columns), _END_COL)
    _write_title_band(ws, row, 1, end_col, title, color)
    _write_caption_band(ws, row + 1, 1, end_col, caption, color)
    header_row = row + 2
    if frame.empty:
        cell = ws.cell(header_row, 1, "None.")
        cell.font = Font(name=FONT_NAME, size=9, italic=True, color=SLATE)
        return header_row + 2
    _write_dataframe_values(ws, frame, header_row, 1)
    _format_header(ws, header_row, 1, len(frame.columns), color, headers=list(frame.columns), amount_columns=frozenset(amounts))
    last = header_row + len(frame)
    _format_body_block(ws, header_row + 1, last, 1, len(frame.columns), SLATE_LIGHT)
    _apply_number_formats(ws, list(frame.columns), header_row + 1, last, 1, amounts, set())
    for r in range(header_row + 1, last + 1):
        for c in range(1, len(frame.columns) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    return last + 2


def build_audit_sheet(wb: Workbook, result: ReconciliationResult) -> None:
    ws = wb.create_sheet(AUDIT_SHEET)
    md = result.metadata
    approved = (result.approval or {}).get("status") == "APPROVED"

    _write_title_band(ws, 1, 1, _END_COL, "AUDIT & CONTROLS", NAVY)
    _write_caption_band(
        ws, 2, 1, _END_COL,
        "What produced this workbook, what was validated when it was generated, and which decisions "
        "(if any) were recorded. A static PASS on this sheet describes the moment of export only; it "
        "does not re-check edits made in Excel afterward.",
        NAVY,
    )

    identity = pd.DataFrame([
        ("Run ID", result.run_id),
        ("Generated (run timestamp)", str(result.run_timestamp)),
        ("Application version", APP_VERSION),
        ("Matching rule version", MATCHING_RULE_VERSION),
        ("Duplicate rule version", DUPLICATE_RULE_VERSION),
        ("Fiscal period / year", f"{md.get('fiscal_period')} / {md.get('fiscal_year')}"),
        ("QuickBooks file", f"{md.get('qb_filename')}  (SHA-256 {md.get('qb_sha256')})"),
        ("Infinium file", f"{md.get('inf_filename')}  (SHA-256 {md.get('inf_sha256')})"),
        ("QuickBooks historical file", (
            f"{md['qb_secondary_filename']}  (SHA-256 {md.get('qb_secondary_sha256')})"
            if md.get("qb_secondary_filename") else "None supplied"
        )),
        ("Infinium historical file", (
            f"{md['inf_secondary_filename']}  (SHA-256 {md.get('inf_secondary_sha256')})"
            if md.get("inf_secondary_filename") else "None supplied"
        )),
        ("Historical rows used / ignored", (
            f"QuickBooks {result.metrics['QuickBooks Secondary Rows Used']} / {result.metrics['QuickBooks Secondary Rows Ignored']};"
            f" Infinium {result.metrics['Infinium Secondary Rows Used']} / {result.metrics['Infinium Secondary Rows Ignored']}"
        )),
        ("Reviewer adjustments applied", f"{len(result.review_adjustments)} (see below)"),
    ], columns=["Item", "Value"])
    row = _write_table(ws, 4, identity, "RUN IDENTITY", "Compare these across exports to tell runs apart.", SLATE, set())

    approval = pd.DataFrame([
        ("Engine proposed JE (never changed)", result.metrics["Proposed JE Amount"]),
        ("Reviewer-adjusted JE (calculated)", adjusted_je_amount(result)),
        ("Approval status", approval_status_text(result)),
    ], columns=["Item", "Value"])
    row = _write_table(
        ws, row, approval, "CALCULATION VS APPROVAL",
        "A calculated amount is not an approval. 'Final approved' appears only where an explicit approval "
        "was recorded in the application for this run.",
        NAVY, {"Value"},
    )

    controls = result.export_controls.copy()
    if not controls.empty:
        controls = controls[["Check", "Basis", "Status", "Scope", "Detail"]]
    row = _write_table(
        ws, row, controls, "EXPORT CONTROLS",
        "Identity checks first, then counts and dollars, run from the result's own row sets before this "
        "workbook was written. Static: they describe the export, not later edits.",
        SLATE, set(),
    )
    if not controls.empty:
        first = row - len(controls) - 1
        for offset, status in enumerate(controls["Status"]):
            cell = ws.cell(first + offset, 3)
            cell.fill = PatternFill("solid", fgColor=GREEN_LIGHT if status == "PASS" else RED_LIGHT)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color=TEXT)

    live = pd.DataFrame({"Controls that stay live in this workbook": list(LIVE_WORKBOOK_CONTROLS)})
    row = _write_table(
        ws, row, live, "LIVE CONTROLS",
        "Everything not listed here is a snapshot. Record consumption and exact-cent agreement are never "
        "re-checked in Excel -- apply decisions in the application and regenerate the workbook.",
        NAVY, set(),
    )

    adjustments = result.review_adjustments
    row = _write_table(
        ws, row, adjustments, "REVIEWER ADJUSTMENTS",
        "One row per reviewer decision, beside the engine's own unchanged classification.",
        SLATE, {"QuickBooks Amount", "Infinium Amount", "Confirmed Amount Difference", "JE Effect"},
    )
    bridge = result.adjustment_bridge
    row = _write_table(
        ws, row, bridge.drop(columns=["Kind"]) if not bridge.empty else bridge, "ADJUSTMENT BRIDGE",
        "Engine proposed JE -> reviewer adjustments -> reviewer-adjusted JE, with the open and non-posting items.",
        NAVY, {"Amount"},
    )

    classes = (
        result.qb_dispositions["Classification Code"].value_counts().rename_axis("Classification Code")
        .reset_index(name="QuickBooks Rows")
    )
    _write_table(
        ws, row, classes, "ENGINE CLASSIFICATIONS",
        "The engine's stable classification of every primary QuickBooks row. Reviewers never edit these.",
        SLATE, set(),
    )
    for col, width in zip("ABCDEF", (46, 44, 14, 28, 60, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"
    _prepare_sheet(ws, landscape=True)
