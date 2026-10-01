"""Workbook finishing: row autofit and heights, ignored-errors XML, run metadata, saving to bytes."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..config import CENTRAL_TIMEZONE
from ..excel_styles import _autofit_workbook_columns
from ..matching import ReconciliationResult
from ..utils import format_central_timestamp


def _autofit_workbook_rows(wb: Workbook) -> None:
    """
    Dynamically calculates and explicitly sets row heights based on text wrapping 
    and column widths. This forces Excel to auto-expand rows even when the 
    workbook opens in Protected View.
    """
    for ws in wb.worksheets:
        # Capture custom column widths to estimate text wrapping constraints
        col_widths = {}
        for col_letter, dim in ws.column_dimensions.items():
            col_widths[col_letter] = dim.width or 15

        # A merged cell wraps across the whole width of its merge, not the
        # width of its first column -- measured against that one column, a
        # long caption or title reads as a dozen wrapped lines and the row
        # balloons to several times its real height.
        merged_widths: dict[tuple[int, int], float] = {}
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_row == merged_range.max_row:
                merged_widths[(merged_range.min_row, merged_range.min_col)] = sum(
                    col_widths.get(get_column_letter(col), 15)
                    for col in range(merged_range.min_col, merged_range.max_col + 1)
                )
        designed_heights = {**_FIXED_ROW_HEIGHTS.get(ws.title, {}), **getattr(ws, "_fixed_row_heights", {})}

        for row in ws.iter_rows():
            max_lines = 1
            row_idx = row[0].row

            # Row 2 is a controlled report caption band. Do not let its long
            # merged-cell text enter the generic wrapping calculation, which
            # can otherwise expand it to several times the requested height.
            if row_idx == 2:
                ws.row_dimensions[row_idx].height = _controlled_row_two_height(ws.title)
                continue

            fixed_height = designed_heights.get(row_idx)
            if fixed_height is not None:
                ws.row_dimensions[row_idx].height = fixed_height
                continue

            for cell in row:
                # A date renders as MM/DD/YYYY (10 characters), not as the
                # 19-character "YYYY-MM-DD HH:MM:SS" str() of the datetime --
                # measuring the latter wraps every date cell and inflates
                # the whole row.
                if isinstance(cell.value, (datetime, date)):
                    text = "00/00/0000"
                else:
                    text = str(cell.value) if cell.value is not None else ""
                # A formula's text is not what the cell displays, and its result
                # is unknown until Excel calculates -- measuring the formula
                # would wrap a one-line figure or link into a tall row.
                if not text or text.startswith("="):
                    continue

                # Estimate character capacity based on column width
                col_width = merged_widths.get((row_idx, cell.column)) or col_widths.get(cell.column_letter, 15)
                # Approximation: ~1.1 to 1.2 chars fit per Excel width unit (10pt font)
                chars_per_line = max(int(col_width * 1.1), 10)

                cell_lines = 0
                for line in text.split("\n"):
                    # Calculate how many times this line will wrap
                    cell_lines += max(1, (len(line) // chars_per_line) + 1)
                
                if cell_lines > max_lines:
                    max_lines = cell_lines

                # Enable wrap_text for cells taking up multiple lines
                if cell_lines > 1:
                    curr_align = cell.alignment
                    if not curr_align or not curr_align.wrap_text:
                        cell.alignment = Alignment(
                            horizontal=curr_align.horizontal if curr_align else "left",
                            vertical=curr_align.vertical if curr_align else "center",
                            wrap_text=True,
                            shrink_to_fit=curr_align.shrink_to_fit if curr_align else False,
                            indent=curr_align.indent if curr_align else 0
                        )

            # Assign row height. (15 points per line is standard padding)
            if max_lines > 1:
                ws.row_dimensions[row_idx].height = max_lines * 15
            elif row_idx in ws.row_dimensions:
                # Let Excel manage the single-line rows natively
                ws.row_dimensions[row_idx].height = None


                

# Rows whose height is a deliberate design value rather than something to be
# measured from their text. Row 1 is a single-line title band; measured
# against only its first (narrow) column it reads as 3-5 wrapped lines and
# balloons to 45-75pt. Legacy Reconciliation row 3 is the one-line legend.
# Builders mark their own designed rows (title bands, KPI cards, legends, the
# control strip) with fix_row_height; this table covers only what they do not.

_FIXED_ROW_HEIGHTS = {
    "Legacy Reconciliation": {1: 27, 3: 18},
    "Exceptions": {1: 27},
}


def _controlled_row_two_height(sheet_title: str) -> int:
    """Fixed row-2 caption-band height per sheet, overriding autofit."""
    if sheet_title in ("Product Aggregate Summary", "Product Aggregates", "Aggregates"):
        return 60
    if sheet_title == "Legacy Reconciliation":
        return 30
    return 15


def _apply_accountant_output_row_heights(wb: Workbook) -> None:
    """Apply the controlled row-two presentation required by the workpaper."""
    for ws in wb.worksheets:
        ws.row_dimensions[2].height = _controlled_row_two_height(ws.title)


def _save_workbook_bytes(
    wb: Workbook,
    *,
    apply_accountant_row_heights: bool = False,
    skip_autofit_titles: frozenset = frozenset(),
    suppress_text_number_warnings: bool = False,
) -> bytes:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    _autofit_workbook_columns(wb, skip_titles=skip_autofit_titles)
    _autofit_workbook_rows(wb)  # Dynamically auto-expand the row heights
    # Apply fixed reporting requirements after autofit so they cannot be
    # overwritten by content-dependent height calculations.
    if apply_accountant_row_heights:
        _apply_accountant_output_row_heights(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    saved = buffer.getvalue()
    return _add_ignored_errors(saved) if suppress_text_number_warnings else saved


# Worksheet-XML children that must come after <ignoredErrors> in the schema.

_AFTER_IGNORED_ERRORS = (
    "smartTags", "drawing", "legacyDrawing", "legacyDrawingHF", "picture",
    "oleObjects", "controls", "webPublishItems", "tableParts", "extLst",
)


_IGNORED_ERRORS_XML = (
    '<ignoredErrors><ignoredError sqref="A1:XFD1048576" numberStoredAsText="1"/></ignoredErrors>'
)


def _add_ignored_errors(xlsx_bytes: bytes) -> bytes:
    """Suppress Excel's green "number stored as text" triangles on every
    worksheet. Invoice, PO, and customer numbers are identifiers, so storing
    them as text is correct -- but openpyxl has no API for the worksheet's
    <ignoredErrors> element, so it is inserted into each sheet's XML at its
    schema-mandated position after the workbook is saved."""
    source = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    output = io.BytesIO()
    boundary = re.compile(r"<(?:%s)[\s>/]" % "|".join(_AFTER_IGNORED_ERRORS))
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", item.filename):
                xml = data.decode("utf-8")
                match = boundary.search(xml, xml.index("</sheetData>"))
                position = match.start() if match else xml.rindex("</worksheet>")
                data = (xml[:position] + _IGNORED_ERRORS_XML + xml[position:]).encode("utf-8")
            target.writestr(item, data)
    return output.getvalue()


def _apply_workbook_run_metadata(wb: Workbook, result: ReconciliationResult) -> None:
    central_timestamp = result.run_timestamp.astimezone(CENTRAL_TIMEZONE)
    excel_timestamp = central_timestamp.replace(tzinfo=None)
    display_timestamp = format_central_timestamp(central_timestamp)
    # Core properties use the controlled Central timestamp. Run ID and the
    # generation time live on the Posting Summary landing page and in the
    # page footer (printed output only, never competing on-screen with the
    # figures) -- every other sheet's caption states only its own substance.
    # Headers stay blank so Excel cannot surface a stale or locale-generated
    # tag on printed sheets.
    wb.properties.created = excel_timestamp
    wb.properties.modified = excel_timestamp
    existing_description = wb.properties.description or ""
    wb.properties.description = (
        f"{existing_description} Generated {display_timestamp}. Run ID: {result.run_id}."
    ).strip()
    footer_text = f"Run ID: {result.run_id}  |  Generated {display_timestamp}"
    for ws in wb.worksheets:
        for section in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
            section.left.text = None
            section.center.text = None
            section.right.text = None
        for section in (ws.oddFooter, ws.evenFooter, ws.firstFooter):
            section.left.text = None
            section.center.text = footer_text
            section.right.text = "Page &P of &N"
