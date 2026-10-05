"""Sheet names plus Excel table, structured-reference, and totals-formula helpers."""

from __future__ import annotations

from typing import Optional

import pandas as pd
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

from ..matching import (
    INF_ID,
    numeric_quantity_sum,
    numeric_sum,
    QB_ID,
    ReconciliationResult,
)
from ..utils import excel_safe


RECONCILIATION_DETAIL_SHEET = "Reconciliation Detail"


UNRESOLVED_EXCEPTIONS_SHEET = "Unresolved Exceptions"


def _fiscal_period_prefix(result: ReconciliationResult) -> str:
    """"FISCAL PERIOD 05 - 2026": the leading words of every main-sheet title,
    so a reviewer holding two runs of the same workbook open at once knows at
    a glance which period each one belongs to."""
    period = result.metadata.get("fiscal_period")
    year = result.metadata.get("fiscal_year")
    if period is None:
        return f"FISCAL PERIOD NOT SELECTED - {int(year)}" if year else "FISCAL PERIOD NOT SELECTED"
    label = f"FISCAL PERIOD {int(period):02d}"
    return f"{label} - {int(year)}" if year else label


def _write_dataframe_values(ws, frame: pd.DataFrame, start_row: int, start_col: int) -> None:
    for col_offset, header in enumerate(frame.columns):
        ws.cell(start_row, start_col + col_offset, excel_safe(str(header)))
    for row_offset, row in enumerate(frame.values.tolist(), 1):
        for col_offset, value in enumerate(row):
            ws.cell(start_row + row_offset, start_col + col_offset, excel_safe(value))


def _escape_structured_ref_component(text: str) -> str:
    """Escape characters with special meaning inside an Excel structured-table reference."""
    escaped = str(text)
    for char in ("'", "#", "[", "]"):
        escaped = escaped.replace(char, f"'{char}")
    return escaped


def _table_column_reference(table_name: str, column_header: str) -> str:
    """Return a proper qualified structured reference, e.g. ``Table1[Amount]``.

    A bare table name (or ``INDEX(TableName,0,N)`` built from one) is not a
    valid Excel reference when written directly as raw formula text -- only
    Excel's own UI auto-converts a typed table name into this bracketed
    structured-reference form. Writing the bracketed form ourselves is what
    makes formulas outside the table (KPI cards, the proposed JE amount)
    actually resolve instead of showing #NAME?.
    """
    return f"{table_name}[{_escape_structured_ref_component(column_header)}]"


def _table_totals_row_formula(column_header: str) -> str:
    """Return the native Excel table totals-row SUM formula for one column.

    Matches exactly what Excel's own UI writes when a table's Total Row is
    enabled and "Sum" is selected: an *unqualified* single-column reference
    (no table name -- it is implicit from the cell's own position in that
    table's totals row) wrapped in SUBTOTAL so the total also respects any
    filter applied to the table.
    """
    return f"SUBTOTAL(109,[{_escape_structured_ref_component(column_header)}])"


def _review_holds_released_expr() -> str:
    """Review Holds items a reviewer released to the JE (structured
    reference, so it resolves from any sheet, not just Unresolved
    Exceptions -- see the Posting Summary journal-entry bridge)."""
    amount_col = _table_column_reference("ReviewHolds", "Amount")
    disposition_col = _table_column_reference("ReviewHolds", "Reviewer Disposition")
    return f'SUMIFS({amount_col},{disposition_col},"Release to JE")'


def _decisions_missing_support_expr() -> str:
    """Count of Review Holds rows with a reviewer decision but no reviewer,
    review date, or comment -- live, so it follows edits made in Excel."""
    disposition = _table_column_reference("ReviewHolds", "Reviewer Disposition")
    reviewer = _table_column_reference("ReviewHolds", "Reviewer")
    date = _table_column_reference("ReviewHolds", "Review Date")
    comment = _table_column_reference("ReviewHolds", "Comment")
    return (
        f'SUMPRODUCT(({disposition}<>"Pending Review")'
        f'*((({reviewer}="")+({date}="")+({comment}=""))>0))'
    )


def _je_support_manual_exclusions_expr(result: "ReconciliationResult") -> str:
    """True-unmatched JE Support items a reviewer manually excluded, with a
    documented reason -- the engine's own Final Disposition for these rows
    stays TRUE_UNMATCHED; this is a separate, additive manual adjustment."""
    amount_col = _table_column_reference("QuickBooksExceptions", result.qb_mapping["amount"])
    disposition_col = _table_column_reference("QuickBooksExceptions", "Reviewer Disposition")
    return f'SUMIFS({amount_col},{disposition_col},"Exclude")'


def _add_exception_table(
    ws,
    *,
    table_name: str,
    headers: list[str],
    header_row: int,
    total_row: int,
    start_col: int,
    total_label: str,
    summed_headers: set[str],
    style_name: str,
) -> dict[str, int]:
    """Create a filterable exception table with protected, dynamic SUM totals."""
    end_col = start_col + len(headers) - 1
    table = Table(
        displayName=table_name,
        ref=(
            f"{get_column_letter(start_col)}{header_row}:"
            f"{get_column_letter(end_col)}{total_row}"
        ),
        totalsRowCount=1,
        totalsRowShown=True,
    )
    table.tableStyleInfo = TableStyleInfo(
        name=style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    formula_columns: dict[str, int] = {}
    table.tableColumns = []
    for offset, header in enumerate(headers, start=1):
        column = TableColumn(id=offset, name=str(header))
        if offset == 1:
            column.totalsRowLabel = total_label
        if header in summed_headers:
            formula_text = _table_totals_row_formula(header)
            column.totalsRowFunction = "sum"
            formula_columns[header] = start_col + offset - 1
            formula_cell = ws.cell(total_row, start_col + offset - 1)
            formula_cell.value = f"={formula_text}"
            formula_cell.protection = Protection(locked=True)
        table.tableColumns.append(column)

    ws.add_table(table)

    # Users may add, remove, classify, and annotate exception rows. The totals
    # row and every other report formula remain locked by worksheet protection.
    for row in range(header_row + 1, total_row):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).protection = Protection(locked=False)

    for col in range(start_col, end_col + 1):
        ws.cell(total_row, col).protection = Protection(locked=True)
    return formula_columns


def _add_plain_data_table(
    ws,
    *,
    table_name: str,
    start_col: int,
    end_col: int,
    header_row: int,
    last_data_row: int,
    style_name: str = "TableStyleMedium2",
) -> None:
    """Register a read-only data block (raw source data, reconciliation
    detail) as a genuine Excel Table (ListObject), so a screen reader
    announces each column's header as the user navigates down through the
    rows instead of requiring a manual title-reading command.

    Deliberately excludes the sheet's own fixed control/source-total row --
    that row is an audit control amount and must stay exactly as written,
    never recast as Excel's native, filter-sensitive SUBTOTAL totals row.
    """
    if last_data_row < header_row:
        return
    table = Table(
        displayName=table_name,
        ref=(
            f"{get_column_letter(start_col)}{header_row}:"
            f"{get_column_letter(end_col)}{last_data_row}"
        ),
    )
    table.tableStyleInfo = TableStyleInfo(
        name=style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _source_totals(frame: pd.DataFrame, mapping: dict[str, Optional[str]]) -> dict[str, float]:
    totals: dict[str, float] = {}
    amount_col = mapping.get("amount")
    quantity_col = mapping.get("quantity")
    if amount_col:
        totals[amount_col] = numeric_sum(frame[amount_col])
    if quantity_col:
        totals[quantity_col] = numeric_quantity_sum(frame[quantity_col])
    return totals


def _duplicate_source_indexes(
    result: ReconciliationResult,
    dataset: str,
) -> set[int]:
    """Return duplicate primary-row indexes, including for pre-2.9 session results."""
    attribute = "duplicate_qb_rows" if dataset == "QuickBooks" else "duplicate_inf_rows"
    stored_indexes = getattr(result, attribute, None)
    if stored_indexes is not None:
        return {int(index) for index in stored_indexes}

    analysis = getattr(result, "duplicate_analysis", pd.DataFrame())
    if analysis.empty or not {"Dataset", "Source Row IDs"}.issubset(analysis.columns):
        return set()
    duplicate_ids: set[str] = set()
    source_rows = analysis.loc[analysis["Dataset"].eq(dataset), "Source Row IDs"]
    for value in source_rows.dropna().astype(str):
        duplicate_ids.update(item.strip() for item in value.split(";") if item.strip())
    frame = result.qb_work if dataset == "QuickBooks" else result.inf_work
    id_column = QB_ID if dataset == "QuickBooks" else INF_ID
    return {
        int(index)
        for index in frame.index
        if str(frame.at[index, id_column]) in duplicate_ids
    }
