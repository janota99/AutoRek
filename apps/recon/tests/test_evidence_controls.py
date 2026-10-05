"""Evidence, link-conflict, reviewer-decision, and export-control regressions.

Positive cases show the evidence rules accept what they should; the negative
cases are the point: ambiguous, weak, or contradictory records must never be
forced into an accepted match, and an accepted match must never be consumed
twice.
"""

from __future__ import annotations

import dataclasses
import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from apps.recon.matching import (
    apply_review_decisions,
    assess_group_links,
    build_reconciliation,
    clean_alphanumeric,
    clean_po,
    identifier_flags,
    identifier_text,
    record_approval,
    reference_strength,
    ReviewDecision,
    ReviewDecisionError,
    verify_result,
)
from apps.recon.matching.core import KEY_INV, KEY_PO
from apps.recon.matching.engine import perform_matching, prepare_working_frame
from apps.recon.matching.evidence import LinkIndex, extract_reference
from apps.recon.matching.labels import (
    CLASS_ACCEPTED_HISTORICAL_MATCH,
    CLASS_ACCEPTED_MATCH,
    CLASS_ACCEPTED_MATCH_ID_DISCREPANCY,
    CLASS_CONFLICTING_LINKS,
    CLASS_CONSUMED_EVIDENCE,
    CLASS_TRUE_UNMATCHED,
    CLASS_WEAK_EVIDENCE,
    DISPOSITION_MATCHED,
    DISPOSITION_REVIEW_HOLD,
    DISPOSITION_TRUE_UNMATCHED,
)
from apps.recon.workpapers import build_primary_workbook


def _qb(po, invoice, amount, memo="", customer="Acme", date="2026-01-05"):
    return {"PO": po, "Invoice": invoice, "Amount": amount, "Qty": 1, "Period": "1",
            "Customer": customer, "Date": date, "Memo": memo}


def _inf(po, invoice, amount, customer="Acme", date="2026-01-05", period="1"):
    return {"PO": po, "Invoice": invoice, "Amount": amount, "Period": period,
            "Customer": customer, "Date": date}


def _run(qb_rows, inf_rows, qb_mapping, inf_mapping, make_metadata, **kwargs):
    return build_reconciliation(
        pd.DataFrame(qb_rows), pd.DataFrame(inf_rows), qb_mapping, inf_mapping, make_metadata(), 2026, **kwargs,
    )


def _ledger(result):
    return result.qb_dispositions.set_index("QBO Row ID")


# ---------------------------------------------------------------------------
# 1. Identifier normalization
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value", [12345, 12345.0, np.int64(12345), np.float64(12345.0), "12345", " 12345 ", "12345.00", "​12345"])
def test_numeric_and_text_identifiers_normalize_to_the_same_comparison_value(value):
    assert clean_alphanumeric(value) == "12345"
    assert clean_po(value) == "12345"


def test_leading_zeros_are_meaningful_and_never_collapsed():
    assert clean_alphanumeric("0105695") == "0105695"
    assert clean_alphanumeric("0105695") != clean_alphanumeric("105695")
    assert clean_po("00123.0") == "00123"          # the artifact goes, the zeros stay


def test_long_numeric_identifiers_never_pass_through_float_arithmetic():
    assert identifier_text(12345678901234567890) == "12345678901234567890"
    assert clean_alphanumeric(12345678901234567890) == "12345678901234567890"


def test_excel_text_prefix_and_formula_wrapper_are_removed_but_nothing_else():
    assert identifier_text("'0105695") == "0105695"
    assert identifier_text('="0105695"') == "0105695"
    assert identifier_text("  AB-12 / 7 ") == "AB-12 / 7"


def test_scientific_notation_is_flagged_not_repaired():
    assert "EXCEL_SCIENTIFIC_NOTATION" in identifier_flags("1.23457E+11")
    assert clean_alphanumeric("1.23457E+11") != "123457000000"
    assert "PUNCTUATION_IGNORED" in identifier_flags("AB-12")
    assert identifier_flags("12345") == []


def test_punctuation_equivalence_is_applied_but_recorded_on_the_match(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("AB-12", "I1", 10.0)], [_inf("AB12", "I1", 10.0)], qb_mapping, inf_mapping, make_metadata)
    (group,) = result.matches
    assert "equal only after ignoring punctuation" in group.identifier_discrepancy
    assert _ledger(result).loc["QB-1", "Classification Code"] == CLASS_ACCEPTED_MATCH_ID_DISCREPANCY


def test_source_values_are_kept_beside_the_normalized_comparison_values(qb_mapping):
    frame = prepare_working_frame(pd.DataFrame([_qb("PO# 0098", "inv 7", 5.0)]), qb_mapping, "QB", 2026)
    assert frame.at[0, "__REC_SRC_PO"] == "PO# 0098" and frame.at[0, "__REC_NORM_PO"] == "0098"
    assert frame.at[0, "__REC_SRC_INV"] == "inv 7" and frame.at[0, "__REC_NORM_INV"] == "INV7"


# ---------------------------------------------------------------------------
# 2. Configured evidence fields, and the evidence findings
# ---------------------------------------------------------------------------

def test_a_blank_po_is_filled_from_a_configured_description_column(qb_mapping, inf_mapping, make_metadata):
    qb_mapping = {**qb_mapping, "po_fallback": ["Memo"]}
    result = _run(
        [_qb("", "", 40.0, memo="Order PO 4455667 shipped")], [_inf("4455667", "", 40.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    (group,) = result.matches
    assert group.method == "PO + Amount"
    assert "Memo (description)" in group.evidence_summary
    assert _ledger(result).loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED


def test_a_description_with_several_reference_candidates_supplies_nothing(qb_mapping, inf_mapping, make_metadata):
    assert extract_reference("see PO 4455667 and PO 7788991") == ("", "ambiguous")
    assert extract_reference("thanks!") == ("", "")
    qb_mapping = {**qb_mapping, "po_fallback": ["Memo"]}
    result = _run(
        [_qb("", "", 40.0, memo="see PO 4455667 and PO 7788991")], [_inf("4455667", "", 40.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert not result.matches
    assert _ledger(result).loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED


def test_a_fallback_never_overrides_a_populated_dedicated_column(qb_mapping, inf_mapping, make_metadata):
    qb_mapping = {**qb_mapping, "po_fallback": ["Memo"]}
    frame = prepare_working_frame(pd.DataFrame([_qb("111111", "", 5.0, memo="PO 999999")]), qb_mapping, "QB", 2026)
    assert frame.at[0, "__REC_NORM_PO"] == "111111"


def test_no_evidence_is_not_confused_with_rejected_evidence(qb_mapping, inf_mapping, make_metadata):
    result = _run(
        [_qb("P-NONE1", "I-NONE1", 5.0), _qb("P-AMT1", "I-AMT1", 100.0)],
        [_inf("P-AMT1", "I-AMT1", 90.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    findings = dict(zip(result.candidates["QuickBooks Row ID"], result.candidates["Evidence Finding"]))
    assert findings["QB-1"] == "NO_EVIDENCE"
    assert findings["QB-2"] == "EVIDENCE_FOUND_AMOUNT_DIFFERS"
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert "differs by" in ledger.loc["QB-2", "Candidate Evidence"] and "INF-1" in ledger.loc["QB-2", "Candidate Evidence"]


def test_evidence_consumed_by_another_match_is_a_hold_not_no_matching_records(qb_mapping, inf_mapping, make_metadata):
    result = _run(
        [_qb("PO7001", "I-1", 100.0), _qb("PO7001", "I-2", 50.0)],
        [_inf("PO7001", "I-1", 100.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    ledger = _ledger(result)
    assert ledger.loc["QB-1", "Final Disposition"] == DISPOSITION_MATCHED
    assert ledger.loc["QB-2", "Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert ledger.loc["QB-2", "Classification Code"] == CLASS_CONSUMED_EVIDENCE
    assert ledger.loc["QB-2", "Evidence Finding"] == "EVIDENCE_ALREADY_CONSUMED"
    assert result.metrics["Proposed JE Amount"] == 0


# ---------------------------------------------------------------------------
# 3-4. Hierarchy, competing and contradictory links
# ---------------------------------------------------------------------------

def _conflict_rows():
    qb = [_qb("1001", "A200", 100.0)]
    inf = [_inf("1001", "B300", 100.0), _inf("2345", "A200", 100.0)]
    return qb, inf


def test_a_po_match_may_not_override_an_invoice_that_points_to_another_transaction(qb_mapping, inf_mapping, make_metadata):
    qb, inf = _conflict_rows()
    result = _run(qb, inf, qb_mapping, inf_mapping, make_metadata)
    assert result.matches == []                                    # neither link wins
    row = _ledger(result).loc["QB-1"]
    assert row["Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert row["Classification Code"] == CLASS_CONFLICTING_LINKS
    assert row["Reason Code"] == "REVIEW_HOLD_CONFLICTING_LINKS"
    assert "INF-1" in row["Candidate Evidence"] and "INF-2" in row["Candidate Evidence"]
    assert result.metrics["Proposed JE Amount"] == 0               # held, so not accrued
    hold = result.reference_hold_analysis.iloc[0]
    assert set(hold["Related Infinium Row IDs"].split("; ")) == {"INF-1", "INF-2"}


def test_a_conflict_outcome_does_not_depend_on_processing_order(qb_mapping, inf_mapping, make_metadata):
    qb, inf = _conflict_rows()
    forward = _run(qb, inf, qb_mapping, inf_mapping, make_metadata)
    reverse = _run(qb, list(reversed(inf)), qb_mapping, inf_mapping, make_metadata)
    assert forward.matches == [] and reverse.matches == []
    assert _ledger(forward).loc["QB-1", "Reason Code"] == _ledger(reverse).loc["QB-1", "Reason Code"]


def test_a_noncompeting_identifier_discrepancy_is_accepted_and_labeled(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("1001", "A200", 100.0)], [_inf("1001", "B300", 100.0)], qb_mapping, inf_mapping, make_metadata)
    (group,) = result.matches
    assert group.method == "PO + Amount"
    assert "Invoice differs" in group.identifier_discrepancy and "not linked to any other record" in group.identifier_discrepancy
    row = _ledger(result).loc["QB-1"]
    assert row["Classification Code"] == CLASS_ACCEPTED_MATCH_ID_DISCREPANCY
    assert "Non-competing discrepancy" in row["Evidence Summary"]
    assert result.match_register.iloc[0]["Identifier Discrepancy"] == group.identifier_discrepancy


def test_assess_group_links_distinguishes_clean_discrepancy_and_conflict(qb_mapping, inf_mapping):
    qb = prepare_working_frame(pd.DataFrame([_qb("1001", "A200", 1.0)]), qb_mapping, "QB", 2026)
    inf = prepare_working_frame(
        pd.DataFrame([_inf("1001", "A200", 1.0), _inf("1001", "B300", 1.0), _inf("2345", "A200", 1.0)]),
        inf_mapping, "INF", 2026,
    )
    q_links, i_links = LinkIndex.from_frame(qb), LinkIndex.from_frame(inf)
    assert assess_group_links(qb, inf, [0], [0], {"PO", "INV"}, q_links, i_links).status == "CLEAN"
    assert assess_group_links(qb, inf, [0], [1], {"PO"}, q_links, i_links).status == "CONFLICT"


def test_the_stronger_agreement_still_wins_when_it_is_the_only_competitor(qb_mapping, inf_mapping, make_metadata):
    # PO+invoice+amount agree on INF-1 -- a genuine strongest pass -- so INF-2 (same amount, other PO) is free.
    result = _run(
        [_qb("1001", "A200", 100.0)],
        [_inf("1001", "A200", 100.0), _inf("2345", "Z999", 100.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    (group,) = result.matches
    assert group.method == "PO + Invoice + Amount" and group.inf_rows == [0]


# ---------------------------------------------------------------------------
# 5. Weak references and reuse
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    ("NA", "WEAK"), ("TBD", "WEAK"), ("0", "WEAK"), ("0000", "WEAK"), ("XXXX", "WEAK"), ("", "NONE"),
    ("4455667", "SPECIFIC"), ("0105695", "SPECIFIC"), ("1", "SPECIFIC"), ("AB12", "SPECIFIC"),
])
def test_reference_strength(value, expected):
    assert reference_strength(clean_po(value)) == expected


def test_a_weak_reference_with_an_exact_amount_is_held_never_accepted(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("NA", "I-1", 100.0)], [_inf("NA", "I-9", 100.0)], qb_mapping, inf_mapping, make_metadata)
    assert result.matches == []
    row = _ledger(result).loc["QB-1"]
    assert row["Final Disposition"] == DISPOSITION_REVIEW_HOLD
    assert row["Classification Code"] == CLASS_WEAK_EVIDENCE
    assert row["Reason Code"] == "REVIEW_HOLD_WEAK_REFERENCE"
    assert result.metrics["Proposed JE Amount"] == 0


def test_a_weak_reference_with_a_different_amount_is_only_an_informational_flag(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("NA", "I-1", 100.0)], [_inf("NA", "I-9", 90.0)], qb_mapping, inf_mapping, make_metadata)
    row = _ledger(result).loc["QB-1"]
    assert row["Final Disposition"] == DISPOSITION_TRUE_UNMATCHED          # still JE support
    assert row["Evidence Finding"] == "EVIDENCE_INSUFFICIENT_WEAK_REFERENCE"
    assert "weak" in row["Evidence Summary"].lower()


def test_a_weak_po_does_not_stop_a_specific_invoice_from_matching(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("NA", "INV5551", 100.0)], [_inf("NA", "INV5551", 100.0)], qb_mapping, inf_mapping, make_metadata)
    (group,) = result.matches
    assert group.method == "Invoice + Amount"


def test_equal_amounts_and_repeated_names_do_not_prove_a_match(qb_mapping, inf_mapping, make_metadata):
    # Same amount and customer, but no shared PO/invoice at all: nothing is matched or held.
    result = _run([_qb("PO-A1", "I-A1", 100.0)], [_inf("PO-B2", "I-B2", 100.0)], qb_mapping, inf_mapping, make_metadata)
    assert result.matches == []
    row = _ledger(result).loc["QB-1"]
    assert row["Final Disposition"] == DISPOSITION_TRUE_UNMATCHED and row["Evidence Finding"] == "NO_EVIDENCE"


def test_a_repeated_business_reference_alone_does_not_exclude_a_duplicate(qb_mapping, inf_mapping, make_metadata):
    result = _run(
        [_qb("PO7777", "I-1", 100.0, customer="Acme", date="2026-01-05"),
         _qb("PO7777", "I-2", 100.0, customer="Beta", date="2026-02-09")],
        [_inf("PO0042", "I-42", 1.0)], qb_mapping, inf_mapping, make_metadata,
    )
    assert result.duplicate_qb_rows == []          # nothing excluded on the reference alone


# ---------------------------------------------------------------------------
# Signed amounts, zero-dollar rows, record consumption
# ---------------------------------------------------------------------------

def test_credits_match_only_at_the_same_signed_amount(qb_mapping, inf_mapping, make_metadata):
    matched = _run([_qb("PO8001", "I-1", -75.25)], [_inf("PO8001", "I-1", -75.25)], qb_mapping, inf_mapping, make_metadata)
    assert len(matched.matches) == 1
    reversed_sign = _run([_qb("PO8001", "I-1", 75.25)], [_inf("PO8001", "I-1", -75.25)], qb_mapping, inf_mapping, make_metadata)
    assert reversed_sign.matches == []                      # no absolute-value matching
    assert _ledger(reversed_sign).loc["QB-1", "Final Disposition"] == DISPOSITION_REVIEW_HOLD


def test_one_cent_is_not_a_match(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("PO8002", "I-1", 100.00)], [_inf("PO8002", "I-1", 100.01)], qb_mapping, inf_mapping, make_metadata)
    assert result.matches == []


def test_zero_dollar_rows_are_handled_consistently(qb_mapping, inf_mapping, make_metadata):
    result = _run([_qb("PO8003", "I-1", 0.0)], [_inf("PO8003", "I-1", 0.0)], qb_mapping, inf_mapping, make_metadata)
    assert len(result.matches) == 1
    assert result.metrics["Proposed JE Amount"] == 0
    assert (verify_result(result)["Status"] == "PASS").all()


def test_grouped_matches_consume_every_member_exactly_once(qb_mapping, inf_mapping, make_metadata):
    result = _run(
        [_qb("PO9001", "I-A", 60.0), _qb("PO9001", "I-B", 40.0), _qb("PO9002", "I-C", 7.0)],
        [_inf("PO9001", "", 100.0), _inf("PO9002", "I-C", 7.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    grouped = [g for g in result.matches if g.group_level]
    assert len(grouped) == 1 and sorted(grouped[0].qb_rows) == [0, 1]
    controls = verify_result(result)
    assert (controls["Status"] == "PASS").all()
    consumed = [i for g in result.matches for i in g.inf_rows]
    assert len(consumed) == len(set(consumed))


# ---------------------------------------------------------------------------
# 7. Historical matching
# ---------------------------------------------------------------------------

def _hist_inputs(qb_mapping, inf_mapping, make_metadata, hist_rows, qb_rows=None):
    return _run(
        qb_rows or [_qb("PO6001", "I-6", 55.0)], [_inf("PO0000", "I-0", 1.0)],
        qb_mapping, inf_mapping, lambda: make_metadata(inf_secondary_filename="inf_prior.xlsx"),
        inf_secondary_raw=pd.DataFrame(hist_rows), inf_secondary_mapping=inf_mapping,
    )


def test_a_historical_match_records_its_file_row_period_date_method_and_discrepancy(
    qb_mapping, inf_mapping, make_metadata,
):
    result = _hist_inputs(
        qb_mapping, inf_mapping, make_metadata,
        [_inf("PO6001", "I-HIST", 55.0, period="12", date="2025-12-20")],
    )
    clearance = result.historical_clearances.iloc[0]
    assert clearance["Secondary Source File"] == "inf_prior.xlsx"
    assert clearance["Secondary Row ID"] == "INF-HIST-1"
    assert clearance["Secondary Period"] == "P12-2026"
    assert clearance["Secondary Date"] == "2025-12-20"
    assert "PO + Amount" in clearance["Match Method"]
    assert "Invoice differs" in clearance["Identifier Discrepancy"]
    assert _ledger(result).loc["QB-1", "Classification Code"] == CLASS_ACCEPTED_HISTORICAL_MATCH


def test_multiple_historical_candidates_are_not_forced_into_a_match(qb_mapping, inf_mapping, make_metadata):
    result = _hist_inputs(
        qb_mapping, inf_mapping, make_metadata,
        [_inf("PO6001", "I-H1", 55.0, period="12"), _inf("PO6001", "I-H2", 55.0, period="12")],
    )
    assert result.historical_clearances.empty
    assert _ledger(result).loc["QB-1", "Final Disposition"] == DISPOSITION_TRUE_UNMATCHED


def test_a_historical_conflict_is_refused_like_a_primary_conflict(qb_mapping, inf_mapping, make_metadata):
    result = _hist_inputs(
        qb_mapping, inf_mapping, make_metadata,
        [_inf("PO6001", "I-H1", 55.0), _inf("PO7002", "I-6", 55.0)],
        qb_rows=[_qb("PO6001", "I-6", 55.0)],
    )
    assert result.historical_clearances.empty


def test_unused_historical_rows_never_create_exceptions(qb_mapping, inf_mapping, make_metadata):
    result = _hist_inputs(
        qb_mapping, inf_mapping, make_metadata,
        [_inf("PO6001", "I-6", 55.0), _inf("PO5555", "I-X", 999.0)],
    )
    assert result.metrics["Infinium Secondary Rows Used"] == 1
    assert result.metrics["Infinium Secondary Rows Ignored"] == 1
    assert "INF-HIST-2" not in set(result.inf_work["__REC_INF_ID"])        # never a primary exception


# ---------------------------------------------------------------------------
# 8. Reviewer decisions
# ---------------------------------------------------------------------------

@pytest.fixture
def reviewable(qb_mapping, inf_mapping, make_metadata):
    """QB-1 matched; QB-2 held (amount differs); QB-3 held (weak, exact amount);
    QB-4 true unmatched; INF-3 is a free Infinium row."""
    return _run(
        [_qb("PO1111", "I-1", 10.0), _qb("PO2222", "I-2", 100.0), _qb("NA", "I-3", 30.0), _qb("PO4444", "I-4", 44.0)],
        [_inf("PO1111", "I-1", 10.0), _inf("PO2222", "I-2", 90.0), _inf("NA", "I-9", 30.0)],
        qb_mapping, inf_mapping, make_metadata,
    )


def _decision(row, action, **overrides):
    values = dict(qb_row_id=row, action=action, reviewer="J. Reviewer", decided_on="2026-02-20", reason="Checked source documents")
    values.update(overrides)
    return ReviewDecision(**values)


def test_reviewer_adjustments_bridge_the_engine_je_without_changing_the_engine(reviewable):
    before = reviewable.qb_dispositions.copy()
    engine_je = reviewable.metrics["Proposed JE Amount"]
    reviewed = apply_review_decisions(reviewable, [
        _decision("QB-2", "RELEASE_TO_JE"),
        _decision("QB-4", "EXCLUDE", support_reference="Credit memo 88"),
        _decision("QB-3", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
    ])
    pd.testing.assert_frame_equal(reviewable.qb_dispositions, before)          # engine untouched
    pd.testing.assert_frame_equal(reviewed.qb_dispositions, before)
    assert reviewed.metrics["Proposed JE Amount"] == engine_je
    bridge = reviewed.adjustment_bridge.set_index("Kind")
    assert bridge.loc["engine", "Amount"] == pytest.approx(44.0)
    assert reviewed.adjustment_bridge.loc[reviewed.adjustment_bridge["Line"].str.startswith("Plus"), "Amount"].iloc[0] == pytest.approx(100.0)
    assert reviewed.adjustment_bridge.loc[reviewed.adjustment_bridge["Line"].str.startswith("Less"), "Amount"].iloc[0] == pytest.approx(-44.0)
    assert bridge.loc["result", "Amount"] == pytest.approx(100.0)                # 44 + 100 - 44
    adjustments = reviewed.review_adjustments.set_index("QuickBooks Row ID")
    assert adjustments.loc["QB-2", "Engine Classification Code"] == "AMOUNT_VARIANCE"
    assert adjustments.loc["QB-3", "Confirmed Amount Difference"] == pytest.approx(0.0)
    assert (verify_result(reviewed)["Status"] == "PASS").all()


def test_decisions_need_reviewer_date_reason_and_the_right_row_state(reviewable):
    for bad, fragment in [
        (_decision("QB-2", "RELEASE_TO_JE", reviewer=""), "reviewer name"),
        (_decision("QB-2", "RELEASE_TO_JE", decided_on="not a date"), "decision date"),
        (_decision("QB-2", "RELEASE_TO_JE", reason=" "), "reason"),
        (_decision("QB-1", "RELEASE_TO_JE"), "only to review holds"),
        (_decision("QB-4", "RELEASE_TO_JE"), "only to review holds"),
        (_decision("QB-4", "EXCLUDE"), "supporting reference"),
        (_decision("QB-99", "EXCLUDE", support_reference="x"), "not a primary QuickBooks row"),
        (_decision("QB-2", "SHRUG"), "unknown action"),
    ]:
        with pytest.raises(ReviewDecisionError, match=fragment):
            apply_review_decisions(reviewable, [bad])


def test_a_confirmed_manual_match_respects_record_consumption(reviewable):
    # INF-1 was consumed by an accepted match (M-001).
    with pytest.raises(ReviewDecisionError, match="already consumed"):
        apply_review_decisions(reviewable, [_decision("QB-2", "CONFIRM_MATCH", inf_row_ids=("INF-1",))])
    with pytest.raises(ReviewDecisionError, match="must name the Infinium"):
        apply_review_decisions(reviewable, [_decision("QB-2", "CONFIRM_MATCH")])
    with pytest.raises(ReviewDecisionError, match="not a primary Infinium row"):
        apply_review_decisions(reviewable, [_decision("QB-2", "CONFIRM_MATCH", inf_row_ids=("INF-77",))])
    # Two decisions may not both claim INF-3.
    with pytest.raises(ReviewDecisionError, match="already claimed"):
        apply_review_decisions(reviewable, [
            _decision("QB-3", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
            _decision("QB-2", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
        ])
    with pytest.raises(ReviewDecisionError, match="more than one decision"):
        apply_review_decisions(reviewable, [_decision("QB-2", "CARRY_FORWARD"), _decision("QB-2", "RELEASE_TO_JE")])


def test_a_confirmed_match_with_a_different_amount_is_recorded_not_posted(reviewable):
    reviewed = apply_review_decisions(reviewable, [_decision("QB-2", "CONFIRM_MATCH", inf_row_ids=("INF-2",))])
    row = reviewed.review_adjustments.iloc[0]
    assert row["Confirmed Amount Difference"] == pytest.approx(10.0) and row["JE Effect"] == 0
    assert reviewed.adjustment_bridge.set_index("Kind").loc["result", "Amount"] == pytest.approx(44.0)


def test_approval_is_separate_from_calculation(reviewable):
    # Open holds block approval.
    with pytest.raises(ReviewDecisionError, match="no resolving decision"):
        record_approval(apply_review_decisions(reviewable, []), "A. Approver", "2026-02-21")
    reviewed = apply_review_decisions(reviewable, [
        _decision("QB-2", "RELEASE_TO_JE"), _decision("QB-3", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
    ])
    assert reviewed.approval["status"] == "NOT APPROVED"
    carried = apply_review_decisions(reviewable, [
        _decision("QB-2", "CARRY_FORWARD"), _decision("QB-3", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
    ])
    with pytest.raises(ReviewDecisionError, match="no resolving decision"):
        record_approval(carried, "A. Approver", "2026-02-21")
    with pytest.raises(ReviewDecisionError, match="approver name"):
        record_approval(reviewed, "", "2026-02-21")
    approved = record_approval(reviewed, "A. Approver", "2026-02-21", "ok")
    assert approved.approval["status"] == "APPROVED"
    assert approved.approval["approved_amount"] == pytest.approx(144.0)
    assert approved.approval["run_id"] == reviewed.run_id


def _sheet_text(workbook_bytes, sheet):
    ws = load_workbook(io.BytesIO(workbook_bytes))[sheet]
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_the_workbook_says_final_approved_only_after_an_approval_is_recorded(reviewable):
    reviewed = apply_review_decisions(reviewable, [
        _decision("QB-2", "RELEASE_TO_JE"), _decision("QB-3", "CONFIRM_MATCH", inf_row_ids=("INF-3",)),
    ])
    assert "FINAL APPROVED JE" not in _sheet_text(build_primary_workbook(reviewed), "Posting Summary")
    assert "NOT APPROVED" in _sheet_text(build_primary_workbook(reviewed), "Posting Summary")
    approved = record_approval(reviewed, "A. Approver", "2026-02-21")
    text = _sheet_text(build_primary_workbook(approved), "Posting Summary")
    assert "FINAL APPROVED JE (recorded approval)" in text and "APPROVED by A. Approver" in text


def test_decisions_open_in_the_workbook_through_the_existing_reviewer_columns(reviewable):
    reviewed = apply_review_decisions(reviewable, [_decision("QB-2", "RELEASE_TO_JE")])
    ws = load_workbook(io.BytesIO(build_primary_workbook(reviewed)))["Unresolved Exceptions"]
    values = {c.value for row in ws.iter_rows() for c in row if c.value is not None}
    assert "Release to JE" in values and "J. Reviewer" in values and "2026-02-20" in values


def test_the_audit_sheet_identifies_the_run_versions_files_and_controls(qb_mapping, inf_mapping, make_metadata):
    result = _hist_inputs(qb_mapping, inf_mapping, make_metadata, [_inf("PO6001", "I-6", 55.0)])
    text = _sheet_text(build_primary_workbook(result), "Audit & Controls")
    from apps.recon.matching import APP_VERSION, MATCHING_RULE_VERSION
    for needle in (result.run_id, APP_VERSION, MATCHING_RULE_VERSION, "inf_prior.xlsx", "qb.xlsx",
                   "Validated at export (static)", "LIVE CONTROLS", "NOT APPROVED"):
        assert needle in text


# ---------------------------------------------------------------------------
# 9. Export controls: identities first, then counts and dollars
# ---------------------------------------------------------------------------

def test_every_export_control_passes_on_a_mixed_result(reviewable):
    controls = verify_result(reviewable)
    assert (controls["Status"] == "PASS").all()
    assert {"Identity", "Count", "Amount"} <= set(controls["Basis"])
    assert (controls["Scope"] == "Validated at export (static)").all()


def _failed(result):
    controls = verify_result(result)
    return set(controls.loc[controls["Status"] == "FAIL", "Check"])


def test_a_row_consumed_twice_fails_even_when_every_total_still_ties(qb_mapping, inf_mapping, make_metadata):
    result = _run(
        [_qb("PO1111", "I-1", 10.0), _qb("PO2222", "I-2", 10.0)],
        [_inf("PO1111", "I-1", 10.0), _inf("PO2222", "I-2", 10.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert not _failed(result)
    twin = dataclasses.replace(result.matches[1], inf_rows=list(result.matches[0].inf_rows))
    tampered = dataclasses.replace(result, matches=[result.matches[0], twin])
    failed = _failed(tampered)
    assert "No source record is consumed by more than one match or clearance" in failed


def test_offsetting_identity_errors_are_caught_although_counts_and_dollars_agree(qb_mapping, inf_mapping, make_metadata):
    # QB-1 (matched) and QB-2 (true unmatched) carry the same amount: swap them between populations.
    result = _run(
        [_qb("PO1111", "I-1", 10.0), _qb("PO2222", "I-2", 10.0)],
        [_inf("PO1111", "I-1", 10.0)],
        qb_mapping, inf_mapping, make_metadata,
    )
    assert not _failed(result)
    swapped_group = dataclasses.replace(result.matches[0], qb_rows=[1])
    tampered = dataclasses.replace(result, matches=[swapped_group], unmatched_qb=[0])
    failed = _failed(tampered)
    assert "Result populations and the disposition ledger contain the same rows" in failed
    # Counts and dollars alone could not see it:
    assert tampered.metrics["Final Disposition - Matched Rows"] == 1 and len(tampered.unmatched_qb) == 1


def test_a_proposed_je_that_disagrees_with_exception_support_fails(reviewable):
    tampered = dataclasses.replace(reviewable, metrics={**reviewable.metrics, "Proposed JE Amount": 45.0})
    assert "Proposed JE amount equals the sum of its exception rows" in _failed(tampered)


def test_a_missing_disposition_row_fails(reviewable):
    tampered = dataclasses.replace(reviewable, qb_dispositions=reviewable.qb_dispositions.iloc[1:].reset_index(drop=True))
    assert "Every primary QuickBooks row has exactly one final disposition" in _failed(tampered)


def test_a_hold_report_that_disagrees_with_its_held_rows_fails(reviewable):
    tampered = dataclasses.replace(reviewable, reference_hold_analysis=reviewable.reference_hold_analysis.iloc[1:])
    assert "Reference-evidence hold report matches its held rows" in _failed(tampered)


def test_a_workbook_is_not_written_when_a_control_fails(reviewable):
    tampered = dataclasses.replace(reviewable, metrics={**reviewable.metrics, "Proposed JE Amount": 45.0})
    with pytest.raises(ValueError, match="Export control failure"):
        build_primary_workbook(tampered)


def test_accepted_matches_always_agree_to_the_cent(reviewable):
    for group in reviewable.matches:
        q = sum(int(reviewable.qb_work.at[i, "__REC_AMOUNT_CENTS"]) for i in group.qb_rows) if "__REC_AMOUNT_CENTS" in reviewable.qb_work else None
        if q is not None:
            i_total = sum(int(reviewable.inf_work.at[i, "__REC_AMOUNT_CENTS"]) for i in group.inf_rows)
            assert q == i_total
