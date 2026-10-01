"""Guards for excel_styles' style reuse.

_CellPainter copies openpyxl's internal ``cell._style`` array instead of
assigning font/fill/border/alignment cell by cell (see its docstring). That
only works while openpyxl keeps a cell's look in that array of style-list
positions. requirements.txt pins openpyxl to 3.1.x; if a future version
changes the internals, these tests fail before a workbook silently loses
its formatting.
"""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from apps.recon.excel_styles import (
    _apply_number_formats,
    _attribute_painter,
    _format_body_block,
)

LOOK = {
    "font": Font(name="Segoe UI", size=10, color="172B4D"),
    "border": Border(bottom=Side(style="thin", color="D5DDE4")),
    "fill": PatternFill("solid", fgColor="F5F7FA"),
    "alignment": Alignment(horizontal="left", vertical="center"),
}


def _seed(ws):
    """Cells with different starting styles: untouched, a number format, a
    protection setting, and an earlier font that the look must replace."""
    ws["A1"] = "plain"
    ws["A2"] = 12.5
    ws["A2"].number_format = "0.00"
    ws["A3"] = "unlocked"
    ws["A3"].protection = Protection(locked=False)
    ws["A4"] = "bold"
    ws["A4"].font = Font(bold=True)
    for row in range(5, 9):
        ws.cell(row, 1, row)
    return [ws.cell(row, 1) for row in range(1, 9)]


def _look(cell):
    return (
        repr(cell.font), repr(cell.fill), repr(cell.border), repr(cell.alignment),
        cell.number_format, repr(cell.protection),
    )


def test_cell_painter_matches_openpyxl_assignment():
    """Painting must leave every cell exactly as plain attribute assignment
    would, keep what the cell already had, and register the same styles."""
    assigned_wb, painted_wb = Workbook(), Workbook()
    assigned, painted = _seed(assigned_wb.active), _seed(painted_wb.active)
    for cell in assigned:
        for name, value in LOOK.items():
            setattr(cell, name, value)
    painter = _attribute_painter(**LOOK)
    for cell in painted:
        painter.paint(cell)

    assert [tuple(c._style) for c in painted] == [tuple(c._style) for c in assigned]
    assert [_look(c) for c in painted] == [_look(c) for c in assigned]
    assert painted[1].number_format == "0.00"
    assert painted[2].protection.locked is False
    collections = ("_fonts", "_fills", "_borders", "_alignments", "_protections", "_number_formats")
    for collection in collections:
        assert list(getattr(painted_wb, collection)) == list(getattr(assigned_wb, collection))

    reloaded = load_workbook(io.BytesIO(_save(painted_wb))).active
    reference = load_workbook(io.BytesIO(_save(assigned_wb))).active
    assert [_look(reloaded.cell(r, 1)) for r in range(1, 9)] == [
        _look(reference.cell(r, 1)) for r in range(1, 9)
    ]

    # Copies, not one shared array: restyling one cell leaves the rest alone.
    painted[5].font = Font(italic=True)
    assert painted[6].font.italic is False


def test_body_block_and_number_formats_style_every_cell():
    """The real helpers, end to end: stripes alternate, each column gets its
    format, and the numeric font swap still reads each cell's own font."""
    wb = Workbook()
    ws = wb.active
    headers = ["Customer", "Amount", "Qty"]
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for row in range(2, 8):
        ws.cell(row, 1, f"C{row}")
        ws.cell(row, 2, row * 1.5)
        ws.cell(row, 3, row)
    _format_body_block(ws, 2, 7, 1, 3, "F5F7FA")
    _apply_number_formats(ws, headers, 2, 7, 1, {"Amount"}, {"Qty"})

    for row in range(2, 8):
        expected_fill = "00F5F7FA" if row % 2 == 0 else None
        for col in range(1, 4):
            cell = ws.cell(row, col)
            assert (cell.fill.fgColor.rgb if cell.fill.fill_type else None) == expected_fill
            assert cell.border.left.style == "thin"
        assert ws.cell(row, 1).font.name == "Segoe UI"
        assert ws.cell(row, 2).font.name == "Consolas" and ws.cell(row, 2).font.size == 9
        assert "$" in ws.cell(row, 2).number_format
        assert ws.cell(row, 3).alignment.horizontal == "right"


def _save(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
