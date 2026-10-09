"""Excel workbook for a custom reconciliation: summary and controls, matches, exceptions, and both
datasets with the Match ID column placed where the user asked. No Streamlit."""

from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import ReconResult
from .spec import NORMALIZERS

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MONEY = '#,##0.00;[Red](#,##0.00)'
_MONEY_WORDS = ("amount", "total", "difference")


def _sheet_name(label: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", label).strip()[:27] or "Data"
    name, n = base, 2
    while name.casefold() in used:
        name = f"{base} {n}"
        n += 1
    used.add(name.casefold())
    return name


def _style(sheet, frame: pd.DataFrame) -> None:
    for cell in sheet[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for index, column in enumerate(frame.columns, start=1):
        letter = get_column_letter(index)
        sample = [len(str(column))] + [len(str(v)) for v in frame[column].head(200)]
        sheet.column_dimensions[letter].width = min(max(sample) + 2, 60)
        if any(word in str(column).lower() for word in _MONEY_WORDS) and pd.api.types.is_numeric_dtype(frame[column]):
            for cell in sheet[letter][1:]:
                cell.number_format = _MONEY


def _settings_frame(result: ReconResult) -> pd.DataFrame:
    spec = result.spec
    rows = [("Mapping name", spec.name), ("Dataset A", spec.label_a), ("Dataset B", spec.label_b),
            ("Amount compared", f"{spec.label_a}: {spec.amount.a}  |  {spec.label_b}: {spec.amount.b} (equal to the cent)")]
    for number, pair in enumerate(spec.align, start=1):
        rows.append((f"Must align {number}", f"{pair.a}  |  {pair.b}  [{NORMALIZERS[pair.normalize]}]"))
    rows += [
        (f"Unique ID column, {spec.label_a}", spec.id_a or "None (row numbers are used)"),
        (f"Unique ID column, {spec.label_b}", spec.id_b or "None (row numbers are used)"),
        (f"Most {spec.label_a} rows summed into one {spec.label_b} row", spec.max_a_per_b),
        (f"Most {spec.label_b} rows summed into one {spec.label_a} row", spec.max_b_per_a),
        ("Match ID column added", "Yes" if spec.add_match_id else "No"),
    ]
    return pd.DataFrame(rows, columns=["Setting", "Value"])


def build_workbook(result: ReconResult) -> bytes:
    used = {"summary", "matches", "match detail", "exceptions"}
    sheets = [
        ("Summary", None), ("Matches", result.matches), ("Match Detail", result.detail),
        ("Exceptions", result.exceptions),
        (_sheet_name(result.spec.label_a, used), result.dataset_a),
        (_sheet_name(result.spec.label_b, used), result.dataset_b),
    ]
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        row = 0
        blocks = [("Overall result: " + result.overall, None), ("Controls", result.controls),
                  ("Dataset bridge", result.bridge), ("Mapping used", _settings_frame(result))]
        summary = writer.book.create_sheet("Summary", 0)
        for title, frame in blocks:
            summary.cell(row=row + 1, column=1, value=title).font = Font(bold=True, size=12)
            row += 1
            if frame is not None:
                frame.to_excel(writer, sheet_name="Summary", startrow=row, index=False)
                for cell in summary[row + 1][: len(frame.columns)]:
                    cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
                for index, column in enumerate(frame.columns, start=1):
                    if any(w in str(column).lower() for w in _MONEY_WORDS) and pd.api.types.is_numeric_dtype(frame[column]):
                        for cell in summary[get_column_letter(index)][row + 1: row + 1 + len(frame)]:
                            cell.number_format = _MONEY
                row += len(frame) + 2
            else:
                row += 1
        for note in result.warnings:
            summary.cell(row=row + 1, column=1, value=note)
            row += 1
        summary.column_dimensions["A"].width = 52
        for letter in "BCDEFGHIJK":
            summary.column_dimensions[letter].width = 22
        for name, frame in sheets[1:]:
            frame.to_excel(writer, sheet_name=name, index=False)
            _style(writer.sheets[name], frame)
        if "Sheet" in writer.book.sheetnames:
            del writer.book["Sheet"]
    return buffer.getvalue()
