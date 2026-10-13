"""The downloadable Master Grid template follows the selected fiscal period."""
import io

import openpyxl

from apps.fifo_inventory.upload_templates import (
    build_master_grid_template, build_receipts_template, build_upload_templates_zip,
)


class _Store:
    def get(self, alias):
        return [{"date": "2026-01-05", "qty": 1000.0, "total_value": 250.0, "unit_cost": 0.25}]


def _first_row(data):
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    return {h.value: c.value for h, c in zip(ws[1], ws[2])}


def test_prior_period_is_prefilled_and_current_quantity_and_value_are_blank():
    row = _first_row(build_master_grid_template(_Store(), 2026, 3))
    assert (row["02"], row["02V"]) == (1000.0, 250.0)       # beginning balance, from the stored layers
    assert row["03"] is None and row["03V"] is None          # the count is typed in; the value is derived
    assert row["04"] == 0 and row["04V"] == 0                # other periods stay empty


def test_period_one_reads_the_prior_years_period_13():
    row = _first_row(build_master_grid_template(_Store(), 2026, 1))
    assert (row["13"], row["13V"]) == (1000.0, 250.0)
    assert row["01"] is None and row["01V"] is None


def test_without_a_period_it_is_the_generic_blank_grid():
    row = _first_row(build_master_grid_template())
    assert all(row[f"{p:02d}"] == 0 for p in range(1, 14))


def test_instructions_no_longer_ask_for_a_current_period_value():
    wb = openpyxl.load_workbook(io.BytesIO(build_master_grid_template(_Store(), 2026, 3)))
    text = " ".join(str(c.value) for c in wb["Instructions"]["A"] if c.value)
    assert "OPTIONAL" in text and "Prepared for FY2026 P03" in text
    assert "AND value columns filled in" not in text and "taken straight from your own spreadsheet" not in text


def test_receipts_example_row_carries_the_selected_period():
    ws = openpyxl.load_workbook(io.BytesIO(build_receipts_template(5))).worksheets[0]
    assert ws.cell(row=2, column=12).value == 5


def test_zip_builds_for_a_period():
    assert build_upload_templates_zip(_Store(), 2026, 3)[:2] == b"PK"
