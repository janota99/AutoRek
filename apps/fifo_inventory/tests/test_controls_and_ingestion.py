"""PASS / REVIEW / FAIL controls, upload parsing and validation, tolerance settings, and the report."""
import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from apps.fifo_inventory import app_settings
from apps.fifo_inventory.excel_export import (
    _compute_product_controls,
    _compute_summary_control_status,
    _parse_layer_date,
    _quantity_is_material,
    _value_is_acceptable,
    _value_is_material,
    export_master_excel,
)
from apps.fifo_inventory.fifo_calculations import calculate_fifo
from apps.fifo_inventory.ingestion import (
    _finite_number,
    extract_period,
    parse_master_grid_upload,
    prepare_receipts_upload,
    validate_master_grid,
)
from apps.fifo_inventory.user_inputs import PRODUCTS

from .conftest import AS_OF, layer, receipts

EMPTY = pd.DataFrame()


def _result(beginning, rows, ending, alias=2, beg_variance_qty=0.0, beg_variance_val=0.0):
    frame = receipts(*rows) if rows else EMPTY
    on_hand, depleted, raw, metrics = calculate_fifo(
        beginning, frame, ending, period=3, fiscal_year=2026, as_of_date=AS_OF)
    return {"alias": alias, "prod_name": f"Product {alias}", "metrics": metrics, "on_hand": on_hand,
            "depleted": depleted, "receivings": raw, "beg_variance_qty": beg_variance_qty,
            "beg_variance_val": beg_variance_val, "alias_found": True, "alias_ambiguous": False,
            "processing_error": None}


# --- materiality is judged at the penny ----------------------------------------------------------

@pytest.mark.parametrize("value, material", [
    (0.0, False), (0.004, False), (0.01, False), (0.011, False), (-0.01, False),
    (0.015, True), (0.02, True), (-0.02, True), (1.0, True),
])
def test_values_are_material_beyond_one_cent_after_rounding(value, material):
    assert _value_is_material(value) is material
    assert _value_is_acceptable(value) is (not material)


def test_quantity_materiality_uses_the_small_computational_tolerance():
    assert not _quantity_is_material(0.00005)
    assert _quantity_is_material(0.001)
    assert _quantity_is_material(-1)


# --- overall status -------------------------------------------------------------------------------

def test_a_clean_period_passes():
    result = _result([layer("2026-01-01", 100, 100.0)], [("2026-03-01", 50, 75.0)], 120)
    overall, checks = _compute_product_controls(result)
    assert overall == "PASS"
    assert {status for _name, status, _msg in checks} == {"PASS"}
    assert _compute_summary_control_status(result)["overall_status"] == "PASS"


def test_a_period_with_no_activity_is_labelled_so():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100)
    assert _compute_product_controls(result)[0] == "PASS - No Activity"
    assert _compute_summary_control_status(result)["overall_status"] == "PASS - No Activity"


def test_a_fully_consumed_product_is_a_zero_balance_pass():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 0)
    assert _compute_summary_control_status(result)["overall_status"] == "PASS - Zero Balance"


def test_unpriced_stock_is_a_review_not_a_fail():
    result = _result([layer("2026-01-01", 100, 0.0)], [], 100)
    status = _compute_summary_control_status(result)
    assert status["overall_status"] == "REVIEW"
    assert "no recorded cost" in status["exception_message"]


def test_a_beginning_quantity_gap_fails_and_says_why():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100, beg_variance_qty=-10)
    status = _compute_summary_control_status(result)
    assert status["overall_status"] == "FAIL"
    assert status["reconciliation_status"] == "FAIL"
    assert "differs from Master Grid by -10" in status["exception_message"]


def test_a_processing_error_fails():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100)
    result["processing_error"] = "boom"
    overall, checks = _compute_product_controls(result)
    assert overall == "FAIL"
    assert ("Processing Integrity", "FAIL", "boom") in checks


def test_a_value_that_does_not_conserve_fails():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100)
    result["metrics"]["value_variance"] = 5.00
    assert _compute_summary_control_status(result)["overall_status"] == "FAIL"


def test_an_ambiguous_or_missing_alias_fails():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100)
    assert _compute_summary_control_status(result, alias_ambiguous=True)["overall_status"] == "FAIL"
    assert _compute_summary_control_status(result, alias_found=False)["overall_status"] == "FAIL"


def test_a_user_tolerance_changes_what_the_report_judges_material():
    result = _result([layer("2026-01-01", 100, 100.0)], [], 100)
    result["metrics"]["value_variance"] = 0.50
    assert _compute_summary_control_status(result)["overall_status"] == "FAIL"
    loose = {"value_tolerance": 1.00}
    assert _compute_summary_control_status(result, tolerances=loose)["overall_status"].startswith("PASS")


# --- layer dates ---------------------------------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("2026-07-21", (2026, 7, 21)),
    ("PD11-26: 07/03/26-07/30/26", (2026, 7, 3)),          # a range label ages from its first date
    ("2026-07-17 (L6-112601)", (2026, 7, 17)),
    ("P10-26: 2026-06-22 to 2026-06-29", (2026, 6, 22)),
    ("Unknown - Legacy Layer", None),
    ("Prior Period Carryover", None),
])
def test_layer_date_labels(text, expected):
    parsed = _parse_layer_date(text)
    assert parsed is None if expected is None else (parsed.year, parsed.month, parsed.day) == expected


# --- numbers and periods -------------------------------------------------------------------------

@pytest.mark.parametrize("raw, expected", [
    ("1,234.50", 1234.5), ("$1,234.50", 1234.5), ("(500)", -500.0), (" 7 ", 7.0), (3, 3.0),
])
def test_accounting_numbers_parse(raw, expected):
    assert _finite_number(raw, "x") == expected


@pytest.mark.parametrize("raw", ["", None, "abc", float("nan"), float("inf"), "n/a"])
def test_blank_text_nan_and_infinity_are_refused(raw):
    with pytest.raises(ValueError):
        _finite_number(raw, "Amount")


@pytest.mark.parametrize("text, period", [
    ("P12", 12), ("p3", 3), ("PD 5", 5), ("Period 7", 7), ("13", 13), ("FY26 P10", 10), ("Period 14", None),
    ("", None), (None, None),
])
def test_period_extraction(text, period):
    assert extract_period(text) == period


def test_a_receipt_row_from_another_fiscal_year_is_not_detected():
    """Known issue 4 in docs/known-issues.md: only the period number is read. Pinned so a fix is deliberate."""
    assert extract_period("P12 FYE 2025") == 12


# --- Master Grid ---------------------------------------------------------------------------------

def _grid(aliases=None, **columns):
    aliases = list(PRODUCTS) if aliases is None else aliases
    data = {"PRODUCT ALIAS": aliases}
    for name, value in columns.items():
        data[name.lstrip("c")] = [value] * len(aliases) if not isinstance(value, list) else value
    return pd.DataFrame(data)


def test_a_valid_master_grid_has_no_errors():
    assert validate_master_grid(_grid(c02=10.0, c03=5.0), 3) == []


def test_master_grid_problems_are_all_reported():
    grid = _grid(aliases=[2, 2, 3], c02=[1.0, 1.0, "x"], c03=[-5.0, 1.0, 1.0])
    errors = " | ".join(validate_master_grid(grid, 3))
    assert "duplicate aliases: [2]" in errors
    assert "missing product aliases" in errors
    assert "nonnumeric or blank quantities" in errors
    assert "negative ending quantities" in errors


def test_master_grid_missing_period_columns_are_named():
    errors = validate_master_grid(_grid(c02=1.0), 3)
    assert any("required period column 03" in e for e in errors)


def test_period_one_requires_period_13_as_its_beginning():
    errors = validate_master_grid(_grid(c01=1.0), 1)
    assert any("period column 13" in e for e in errors)


def test_grid_upload_parsing_flags_negatives_blanks_and_duplicates():
    upload = pd.DataFrame({"Alias": [2, 2, 999], "02": [10, 11, 5], "03": [-1, 4, 6]})
    updates, period_cols, alias_col, unmatched, duplicates, invalid, _values = parse_master_grid_upload(upload)
    assert alias_col == "Alias" and set(period_cols) == {2, 3}
    assert duplicates == [2]
    assert unmatched == [999]
    assert {"alias": 2, "period": 3, "field": "qty", "value": -1} in invalid
    assert updates[2][2]["qty"] == 11                          # the later duplicate row wins the dict, and was flagged


# --- receipts upload -----------------------------------------------------------------------------

def _upload(**overrides):
    row = {"PRODUCT ALIAS": 2, "PERIOD": "P03", "DATE DELIVERED": "2026-03-05",
           "QUANTITY": 100, "PRICE": 250.0, "PO #": "PO-1"}
    row.update(overrides)
    return row


def test_receipts_are_filtered_to_the_selected_period():
    df = pd.DataFrame([_upload(), _upload(PERIOD="P04", **{"PO #": "PO-2"})])
    valid, errors, warnings, _debug = prepare_receipts_upload(df, 3)
    assert errors == []
    assert valid["PO #"].tolist() == ["PO-1"]


def test_rejected_rows_are_reported_with_their_source_rows_and_not_used():
    df = pd.DataFrame([_upload(), _upload(QUANTITY=0), _upload(PRICE=-5), _upload(**{"PRODUCT ALIAS": 999}),
                       _upload(**{"DATE DELIVERED": "nope"})])
    valid, errors, _w, debug = prepare_receipts_upload(df, 3)
    assert len(valid) == 1
    assert "4 selected-period receipt row(s) were rejected" in errors[0]
    assert "[3, 4, 5, 6]" in errors[0]
    assert len(debug["rejected_rows"]) == 4


def test_likely_duplicate_receipts_warn_but_are_kept():
    df = pd.DataFrame([_upload(), _upload()])
    valid, errors, warnings, _d = prepare_receipts_upload(df, 3)
    assert len(valid) == 2 and errors == []
    assert any("may be duplicates" in w for w in warnings)


def test_zero_value_receipts_warn():
    valid, _e, warnings, _d = prepare_receipts_upload(pd.DataFrame([_upload(PRICE=0.0)]), 3)
    assert len(valid) == 1
    assert any("zero total value" in w for w in warnings)


def test_missing_required_fields_are_named_and_nothing_is_returned():
    df = pd.DataFrame([{"PRODUCT ALIAS": 2, "PERIOD": "P03", "QUANTITY": 1, "PRICE": 1.0}])
    valid, errors, _w, _d = prepare_receipts_upload(df, 3)
    assert valid.empty
    assert "Delivery Date" in errors[0]


def test_no_period_column_applies_every_valid_row_and_says_so():
    df = pd.DataFrame([_upload()]).drop(columns=["PERIOD"])
    valid, errors, warnings, _d = prepare_receipts_upload(df, 3)
    assert len(valid) == 1 and errors == []
    assert any("No period column" in w for w in warnings)


def test_receipt_text_amounts_are_read_as_accounting_numbers():
    df = pd.DataFrame([_upload(QUANTITY="1,000", PRICE="$2,500.00")])
    valid, errors, _w, _d = prepare_receipts_upload(df, 3)
    assert errors == []
    assert valid.iloc[0]["QUANTITY"] == 1000 and valid.iloc[0]["PRICE"] == 2500.0


# --- tolerance settings --------------------------------------------------------------------------

def test_defaults_apply_when_no_settings_file_exists(settings_path):
    assert app_settings.load_settings() == app_settings.DEFAULT_TOLERANCES
    assert not app_settings.is_overridden()


def test_saved_settings_take_effect_and_reset_removes_them(settings_path):
    saved = app_settings.save_settings({"qty_tolerance": 0.5, "value_tolerance": 0.05, "zero_cost_tolerance": 0.0})
    assert app_settings.get_effective_tolerances() == saved
    assert app_settings.is_overridden()
    app_settings.reset_to_defaults()
    assert not settings_path.exists() and not app_settings.is_overridden()


@pytest.mark.parametrize("bad", [
    {"qty_tolerance": -1, "value_tolerance": 0.01, "zero_cost_tolerance": 0.0},
    {"qty_tolerance": float("nan"), "value_tolerance": 0.01, "zero_cost_tolerance": 0.0},
    {"qty_tolerance": "1", "value_tolerance": 0.01, "zero_cost_tolerance": 0.0},
    {"qty_tolerance": 1, "value_tolerance": 0.01},
])
def test_invalid_settings_are_refused_and_leave_nothing_behind(settings_path, bad):
    with pytest.raises(ValueError):
        app_settings.save_settings(bad)
    assert not settings_path.exists()


def test_a_corrupt_settings_file_falls_back_to_defaults(settings_path):
    settings_path.write_text("{not json", encoding="utf-8")
    assert app_settings.load_settings() == app_settings.DEFAULT_TOLERANCES
    settings_path.write_text(json.dumps({"qty_tolerance": -5}), encoding="utf-8")
    assert app_settings.load_settings() == app_settings.DEFAULT_TOLERANCES


# --- the Excel report ----------------------------------------------------------------------------

def test_the_master_report_has_a_summary_and_one_sheet_per_product():
    results = [_result([layer("2026-01-01", 100, 100.0)], [("2026-03-01", 50, 75.0)], 120, alias=a) for a in (3, 2)]
    data = export_master_excel(results, 3, fiscal_year=2026)
    wb = load_workbook(io.BytesIO(data.getvalue() if hasattr(data, "getvalue") else data))
    assert len(wb.sheetnames) >= 3
    names = " ".join(wb.sheetnames)
    assert "Summary" in names
    assert all(ws.sheet_view.showGridLines is False for ws in wb.worksheets)
    assert len({n.lower() for n in wb.sheetnames}) == len(wb.sheetnames)
    assert all(len(n) <= 31 for n in wb.sheetnames)
