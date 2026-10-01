"""Excel output for Transaction Cleanup: sheet naming, grouped sheets, and the formatted workbooks."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .cleanup import (
    AMOUNT_COL_IDX,
    clean_key_series,
    LEFT_ALIGN_COL_IDX,
    MAIN_SHEET_NAME,
    NEW_VENDORS_LABEL,
)


# --- Worksheet naming (mirrors VBA CleanWorksheetName / GetUniqueSheetName) ---

_INVALID_SHEET_CHARS = ["\\", "/", ":", "*", "?", '"', "<", ">", "|", "[", "]"]


def clean_worksheet_name(name: str) -> str:
    cleaned = str(name).strip()
    for ch in _INVALID_SHEET_CHARS:
        cleaned = cleaned.replace(ch, "-")
    cleaned = cleaned.strip("'").strip()
    return cleaned or "Uncategorized"


def unique_sheet_name(requested: str, existing_names: set, protected: str) -> str:
    base = clean_worksheet_name(requested)[:31]
    candidate = base
    if candidate.upper() != protected.upper() and candidate not in existing_names:
        return candidate
    suffix_num = 2
    while True:
        suffix = f" ({suffix_num})"
        candidate = base[: 31 - len(suffix)] + suffix
        if candidate.upper() != protected.upper() and candidate not in existing_names:
            return candidate
        suffix_num += 1


def group_output_sheets(result_df: pd.DataFrame, new_vendor_flags: pd.Series):
    """Returns an ordered dict of {sheet_label: sub_dataframe}. Taxability
    groups use the cleaned (upper, space-stripped) value as the label,
    matching the original macro. 'New Vendors' is appended last if any
    rows qualify."""
    groups = {}
    taxability_key = clean_key_series(result_df["Taxability"])
    for key in [k for k in taxability_key.unique() if k]:
        groups[key] = result_df[taxability_key == key].reset_index(drop=True)

    if new_vendor_flags.any():
        groups[NEW_VENDORS_LABEL] = result_df[new_vendor_flags].reset_index(drop=True)

    return groups


# --- Excel output ---

_HEADER_FONT = Font(bold=True)


_HEADER_BORDER = Border(bottom=Side(style="thin"))


_TOTAL_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")


_TOTAL_BORDER = Border(bottom=Side(style="double"))


_CURRENCY_FMT = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'


_TOTAL_CURRENCY_FMT = "$#,##0.00;($#,##0.00);$-"


def write_formatted_sheet(ws, df: pd.DataFrame, amount_col_idx: int = AMOUNT_COL_IDX,
                           left_align_idx=LEFT_ALIGN_COL_IDX, add_total: bool = True):
    """Writes df into ws with formatting equivalent to the VBA
    FormatHeaderRow / totals-row logic: bold centered header with a bottom
    border, autofilter, currency format + left alignment on the relevant
    columns, an optional SUM total row, approximate column autofit, and a
    frozen header row."""
    n_rows, n_cols = df.shape

    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = _HEADER_BORDER

    for row in df.itertuples(index=False):
        ws.append(list(row))

    last_col_letter = get_column_letter(n_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{n_rows + 1}"

    if n_rows > 0 and 0 <= amount_col_idx < n_cols:
        amount_letter = get_column_letter(amount_col_idx + 1)
        for r in range(2, n_rows + 2):
            ws[f"{amount_letter}{r}"].number_format = _CURRENCY_FMT

    for idx in left_align_idx:
        if 0 <= idx < n_cols and n_rows > 0:
            letter = get_column_letter(idx + 1)
            for r in range(2, n_rows + 2):
                ws[f"{letter}{r}"].alignment = Alignment(horizontal="left")

    if add_total and n_rows > 0 and 0 <= amount_col_idx < n_cols:
        total_row = n_rows + 2
        amount_letter = get_column_letter(amount_col_idx + 1)
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=total_row, column=c)
            cell.fill = _TOTAL_FILL
            cell.font = Font(bold=True)
        ws.cell(row=total_row, column=6, value="Total")
        total_cell = ws.cell(
            row=total_row, column=amount_col_idx + 1,
            value=f"=SUM({amount_letter}2:{amount_letter}{n_rows + 1})",
        )
        total_cell.number_format = _TOTAL_CURRENCY_FMT
        total_cell.border = _TOTAL_BORDER

    for c in range(1, n_cols + 1):
        letter = get_column_letter(c)
        max_len = max(
            [len(str(df.columns[c - 1]))]
            + [len(str(v)) for v in df.iloc[:, c - 1].astype(str).tolist()[:500]]
        )
        ws.column_dimensions[letter].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"


def build_workbook(result_df: pd.DataFrame, new_vendor_flags: pd.Series,
                    removed_df: pd.DataFrame = None,
                    main_sheet_name: str = MAIN_SHEET_NAME):
    """Builds the full multi-sheet workbook: main sheet (no total row), one
    grouped sheet per Taxability value, a New Vendors sheet, and (if any
    rows were removed) a Removed Transactions sheet with the original row
    data plus removal reason - so the cleanup can be fully audited, not
    just trusted from a summary count."""
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = main_sheet_name
    write_formatted_sheet(ws_main, result_df, add_total=False)

    groups = group_output_sheets(result_df, new_vendor_flags)
    existing_names = {main_sheet_name}
    sheet_count = 0

    for label, sub_df in groups.items():
        sheet_name = unique_sheet_name(label, existing_names, main_sheet_name)
        existing_names.add(sheet_name)
        ws = wb.create_sheet(sheet_name)
        write_formatted_sheet(ws, sub_df)
        sheet_count += 1

    if removed_df is not None and len(removed_df) > 0:
        sheet_name = unique_sheet_name("Removed Transactions", existing_names, main_sheet_name)
        existing_names.add(sheet_name)
        ws = wb.create_sheet(sheet_name)
        # Removed Transactions has one extra trailing column (Removal
        # Reason) beyond the standard 11-column layout, but Amount stays at
        # AMOUNT_COL_IDX, so the shared formatter still totals it correctly.
        write_formatted_sheet(ws, removed_df, amount_col_idx=AMOUNT_COL_IDX, left_align_idx=LEFT_ALIGN_COL_IDX)
        sheet_count += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, sheet_count


def build_updated_mapping_workbook(df_mapping: pd.DataFrame, new_vendor_classifications: pd.DataFrame) -> BytesIO:
    """Appends newly classified vendors onto the original mapping file, in
    the same column layout process_transactions expects (column 0 = vendor
    key, column 3 = Taxability, column 4 = Grouping - columns 1/2 are left
    blank for new rows since process_transactions never reads them).

    new_vendor_classifications must have "Vendor", "Taxability", "Grouping"
    columns (the shape produced by the "Classify New Vendors" editor)."""
    mapping_cols = list(df_mapping.columns)

    new_rows = []
    for _, row in new_vendor_classifications.iterrows():
        new_row = {c: "" for c in mapping_cols}
        new_row[mapping_cols[0]] = row["Vendor"]
        new_row[mapping_cols[3]] = row["Taxability"]
        new_row[mapping_cols[4]] = row["Grouping"]
        new_rows.append(new_row)

    combined = pd.concat(
        [df_mapping, pd.DataFrame(new_rows, columns=mapping_cols)],
        ignore_index=True,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor Mapping"
    ws.append(mapping_cols)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.border = _HEADER_BORDER
    for row in combined.itertuples(index=False):
        ws.append(list(row))

    last_col_letter = get_column_letter(len(mapping_cols))
    ws.auto_filter.ref = f"A1:{last_col_letter}{len(combined) + 1}"
    ws.freeze_panes = "A2"
    for i, col_name in enumerate(mapping_cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(str(col_name)) + 4)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
